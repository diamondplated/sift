"""Fixture builders.

Almost everything is generated rather than committed: the pathological cases are easier to read as
code than as bytes, and a generated corpus keeps the repo clean and the suite fast. Only a couple
of genuinely hard-to-reproduce files live in tests/data/.
"""
from __future__ import annotations

import json
import os
import time
import uuid
from pathlib import Path

# A stable timestamp — Date.now()-style nondeterminism in fixtures makes failures unreproducible.
FIXED_MS = 1_770_000_000_000


def make_csv(
    directory,
    name: str = "sales.csv",
    rows: int = 1000,
    delim: str = ",",
    crlf: bool = False,
    bom: bool = False,
    preamble: int = 0,
    header: bool = True,
    ragged: bool = False,
    bad_int_row: int | None = None,
    quote_notes: bool = False,
    quoted_newline_row: int | None = None,
    nulls_every: int | None = None,
    empties_every: int | None = None,
    nullish_every: int | None = None,
) -> str:
    """Write a CSV and return its path.

    The knobs map to the failure modes this tool exists to survive: a value that won't cast, a
    quoted field containing a newline (which breaks byte-sample row estimation), a BOM, junk
    preamble lines above the header, and ragged rows.
    """
    path = os.path.join(str(directory), name)
    eol = "\r\n" if crlf else "\n"
    regions = ["West", "Midwest", "South", "Northeast"]
    with open(path, "w", encoding="utf-8", newline="") as f:
        if bom:
            f.write("﻿")
        for i in range(preamble):
            f.write(f"# generated file, junk line {i}{eol}")
        if header:
            f.write(delim.join(["order_id", "region", "amount", "note"]) + eol)
        for i in range(rows):
            region = regions[i % len(regions)]
            if nulls_every and i % nulls_every == 0:
                region = ""  # an unquoted empty field reads as NULL in DuckDB
            amount = f"{i}.50"
            if bad_int_row is not None and i == bad_int_row:
                amount = "N/A"
            note = f"note {i}"
            if empties_every and i % empties_every == 0:
                note = '""'  # explicitly quoted empty string, distinct from NULL
            if nullish_every and i % nullish_every == 0:
                note = "N/A"
            if quote_notes:
                # Real CSV writers quote a whole column, not one cell — which is what makes the
                # estimator's quote sniffing work in practice.
                note = f'"{note}"'
            if quoted_newline_row is not None and i == quoted_newline_row:
                note = '"line one' + eol + 'line two"'
            fields = [str(i), region, amount, note]
            if ragged and i % 97 == 0:
                fields = fields[:2]
            f.write(delim.join(fields) + eol)
    return path


def make_parquet(con, directory, name: str = "t.parquet", rows: int = 1000) -> str:
    path = os.path.join(str(directory), name)
    con.execute(
        f"COPY (SELECT range AS order_id, "
        f"['West','Midwest','South','Northeast'][(range % 4) + 1] AS region, "
        f"(range * 1.5)::DECIMAL(12,2) AS amount, "
        f"CASE WHEN range % 50 = 0 THEN NULL ELSE 'note ' || range END AS note "
        f"FROM range({int(rows)})) TO '{path}' (FORMAT parquet)"
    )
    return path


def make_hive_parquet(con, directory, name: str = "events", rows: int = 400) -> str:
    """A hive-partitioned dataset: dt=.../region=.../part.parquet."""
    root = os.path.join(str(directory), name)
    for d, region in (("2026-08-01", "West"), ("2026-08-01", "South"),
                      ("2026-08-02", "West"), ("2026-08-02", "South")):
        part = os.path.join(root, f"dt={d}", f"region={region}")
        os.makedirs(part, exist_ok=True)
        con.execute(
            f"COPY (SELECT range AS id, (range * 2)::BIGINT AS qty FROM range({rows // 4})) "
            f"TO '{os.path.join(part, 'part-0.parquet')}' (FORMAT parquet)"
        )
    return root


def make_delta(con, directory, name: str = "dtable",
               kept: int = 100, tombstoned: int = 50) -> str:
    """A minimal but real Delta table whose version 1 tombstones a still-present file.

    This is the fixture that makes the Delta test meaningful: a raw parquet glob returns
    kept+tombstoned rows, while delta_scan must return only `kept`. If those two numbers are ever
    equal, the fixture is broken and the test proves nothing.
    """
    root = os.path.join(str(directory), name)
    os.makedirs(os.path.join(root, "_delta_log"), exist_ok=True)
    con.execute(
        f"COPY (SELECT range AS id, 'a' AS g FROM range({kept})) "
        f"TO '{os.path.join(root, 'part-0.parquet')}' (FORMAT parquet)"
    )
    con.execute(
        f"COPY (SELECT range AS id, 'b' AS g FROM range({kept}, {kept + tombstoned})) "
        f"TO '{os.path.join(root, 'part-1.parquet')}' (FORMAT parquet)"
    )
    schema = json.dumps({"type": "struct", "fields": [
        {"name": "id", "type": "long", "nullable": True, "metadata": {}},
        {"name": "g", "type": "string", "nullable": True, "metadata": {}}]})

    def size(p):
        return os.path.getsize(os.path.join(root, p))

    log = os.path.join(root, "_delta_log")
    with open(os.path.join(log, "00000000000000000000.json"), "w") as f:
        f.write(json.dumps({"protocol": {"minReaderVersion": 1, "minWriterVersion": 2}}) + "\n")
        f.write(json.dumps({"metaData": {
            "id": str(uuid.uuid4()),
            "format": {"provider": "parquet", "options": {}},
            "schemaString": schema, "partitionColumns": [],
            "configuration": {}, "createdTime": FIXED_MS}}) + "\n")
        for p in ("part-0.parquet", "part-1.parquet"):
            f.write(json.dumps({"add": {"path": p, "partitionValues": {}, "size": size(p),
                                        "modificationTime": FIXED_MS, "dataChange": True}}) + "\n")
    with open(os.path.join(log, "00000000000000000001.json"), "w") as f:
        f.write(json.dumps({"remove": {
            "path": "part-1.parquet", "deletionTimestamp": FIXED_MS + 1000,
            "dataChange": True, "partitionValues": {},
            "size": size("part-1.parquet")}}) + "\n")
    return root


def make_xlsx(directory, name: str = "book.xlsx") -> str:
    """A workbook with a small sheet, a bigger sheet, and an empty one."""
    import openpyxl

    path = os.path.join(str(directory), name)
    wb = openpyxl.Workbook()
    s = wb.active
    s.title = "Summary"
    s.append(["metric", "value"])
    s.append(["total", 42])
    big = wb.create_sheet("By Store")
    big.append(["store", "sales"])
    for i in range(50):
        big.append([i, i * 3])
    wb.create_sheet("Empty")
    wb.save(path)
    return path


def make_ndjson(directory, name: str = "events.ndjson", rows: int = 200) -> str:
    path = os.path.join(str(directory), name)
    with open(path, "w") as f:
        for i in range(rows):
            f.write(json.dumps({"id": i, "region": ["West", "South"][i % 2],
                                "nested": {"a": i}}) + "\n")
    return path


def make_gzip_csv(directory, name: str = "sales.csv.gz", rows: int = 500) -> str:
    import gzip

    path = os.path.join(str(directory), name)
    with gzip.open(path, "wt", newline="\n") as f:
        f.write("id,region\n")
        for i in range(rows):
            f.write(f"{i},{'West' if i % 2 else 'South'}\n")
    return path


def make_fake_xls(directory, name: str = "legacy.xls") -> str:
    """OLE2 magic bytes only — enough to exercise the legacy-.xls refusal path."""
    path = os.path.join(str(directory), name)
    with open(path, "wb") as f:
        f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512)
    return path
